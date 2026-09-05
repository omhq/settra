import io
import unittest

from unittest.mock import MagicMock, patch

from app.sync import loader
from app.sync.config import default_sync_config, validate_sync_config
from app.sync.inspection import (
    GoogleDriveFile,
    TabularFileInspector,
    TabularInspection,
    display_delimiter,
)


class TabularFileInspectorTests(unittest.TestCase):
    def setUp(self):
        self.inspector = TabularFileInspector()

    def test_detects_supported_formats_from_drive_metadata_and_extension(self):
        cases = [
            (
                GoogleDriveFile(
                    "1",
                    "Sales",
                    "application/vnd.google-apps.spreadsheet",
                ),
                "google_sheets",
            ),
            (GoogleDriveFile("2", "sales.csv", "text/csv"), "csv"),
            (GoogleDriveFile("2b", "sales.csv", "application/vnd.ms-excel"), "csv"),
            (
                GoogleDriveFile("3", "sales.xlsx", "application/octet-stream"),
                "excel",
            ),
            (
                GoogleDriveFile("4", "sales.parquet", "application/octet-stream"),
                "parquet",
            ),
        ]

        for file, expected in cases:
            with self.subTest(name=file.name):
                self.assertEqual(expected, self.inspector.detect_format(file))

    def test_detects_csv_encoding_delimiter_and_leading_header_row(self):
        content = (
            "Sales export;;\r\n"
            ";;\r\n"
            "Order ID;Customer;Amount\r\n"
            "1001;Ada;10.25\r\n"
            "1002;Grace;12.50\r\n"
        ).encode("utf-8")

        encoding, delimiter, rows, header_row = self.inspector.inspect_delimited(
            content,
            file_name="sales.csv",
            parsing={
                "encoding": "auto",
                "delimiter": "auto",
                "header_row": "auto",
            },
        )

        self.assertEqual("utf-8", encoding)
        self.assertEqual(";", delimiter)
        self.assertEqual(3, header_row)
        self.assertEqual("Order ID", rows[header_row - 1][0])

    def test_auto_detection_scans_up_to_one_hundred_rows(self):
        rows = [["Executive summary"], *([[]] * 48)]
        rows.extend(
            [
                ["Order ID", "Customer", "Amount"],
                [1001, "Ada", 10.25],
                [1002, "Grace", 12.5],
            ]
        )

        self.assertEqual(100, self.inspector.SAMPLE_ROWS)
        self.assertEqual(50, self.inspector.detect_header_row(rows))

    def test_respects_explicit_csv_overrides(self):
        content = "ignore\nname|amount\nAda|10\n".encode("cp1252")

        encoding, delimiter, _rows, header_row = self.inspector.inspect_delimited(
            content,
            file_name="sales.txt",
            parsing={
                "encoding": "cp1252",
                "delimiter": "pipe",
                "header_row": 2,
            },
        )

        self.assertEqual("cp1252", encoding)
        self.assertEqual("|", delimiter)
        self.assertEqual("|", display_delimiter(delimiter))
        self.assertEqual(2, header_row)

    def test_rejects_known_non_tabular_drive_files(self):
        file = GoogleDriveFile("pdf-1", "report.pdf", "application/pdf")

        with self.assertRaisesRegex(ValueError, "Unsupported Google Drive file type"):
            self.inspector.detect_format(file, content=b"%PDF-1.7\n")


class TabularParserTests(unittest.TestCase):
    @staticmethod
    def _excel_file() -> GoogleDriveFile:
        return GoogleDriveFile(
            "excel-1",
            "sales.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @staticmethod
    def _workbook_bytes(configure) -> bytes:
        import openpyxl

        workbook = openpyxl.Workbook()
        configure(workbook)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    @staticmethod
    def _excel_config(
        *,
        sheets: list[str] | None = None,
        header_row: str | int = "auto",
        table_rules: dict | None = None,
    ) -> dict:
        config = default_sync_config(
            slug="sales",
            file_id="excel-1",
            file_name="sales.xlsx",
            mime_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            sheets=sheets or ["Orders"],
        )
        config["source"]["format"] = "excel"
        config["source"]["parsing"]["header_row"] = header_row
        config["schema"]["tables"] = table_rules or {}
        return validate_sync_config(config, expected_slug="sales")

    def _extract_excel(self, content: bytes, config: dict):
        inspection = TabularInspection(file=self._excel_file(), format="excel")
        tables = loader._extract_excel(
            config,
            self._excel_file(),
            content,
            TabularFileInspector(),
            inspection,
            set(),
        )
        return tables, inspection

    @staticmethod
    def _discover_file_schema(
        file: GoogleDriveFile,
        source_format: str,
        content: bytes,
    ) -> dict:
        with (
            patch("googleapiclient.discovery.build"),
            patch.object(
                loader,
                "_prepare_google_drive_file",
                return_value=(file, source_format, content),
            ),
        ):
            return loader.discover_google_drive_worksheets(file.file_id, object())

    def test_csv_parser_normalizes_detected_header_and_rows(self):
        config = validate_sync_config(
            default_sync_config(slug="orders", file_id="csv-1"),
            expected_slug="orders",
        )
        file = GoogleDriveFile("csv-1", "orders.csv", "text/csv")
        inspection = TabularInspection(file=file, format="csv")

        tables = loader._extract_csv(
            config,
            file,
            b"Report generated today,,\nOrder ID,Amount,Active\n1,10.5,true\n",
            TabularFileInspector(),
            inspection,
            set(),
        )

        self.assertEqual("orders", tables[0]["table_name"])
        self.assertEqual(2, inspection.parsing["header_row"])
        self.assertEqual(
            {"order_id": "1", "amount": "10.5", "active": "true"},
            tables[0]["rows"][0],
        )

    def test_csv_discovery_returns_row_identity_columns(self):
        discovery = self._discover_file_schema(
            GoogleDriveFile("csv-1", "orders.csv", "text/csv"),
            "csv",
            b"Report generated today,,\nOrder ID,Amount,Active\n1,10.5,true\n",
        )

        self.assertEqual([], discovery["worksheets"])
        self.assertEqual(
            [
                {
                    "name": "orders",
                    "header_row": 2,
                    "columns": ["Order ID", "Amount", "Active"],
                    "columns_truncated": False,
                }
            ],
            discovery["worksheet_schemas"],
        )

    def test_composite_row_key_is_validated_and_mapped_to_loaded_columns(self):
        config = default_sync_config(slug="orders", file_id="sheet-1")
        config["schema"]["tables"] = {
            "Orders": {
                "row_key": {
                    "columns": ["Account ID", "Order ID"],
                    "format": "ORD-{Account ID}-{Order ID}",
                },
                "columns": {"Order ID": {"name": "order_number"}},
            }
        }
        config = validate_sync_config(config, expected_slug="orders")

        table = loader._record_table(
            config,
            source_name="Orders",
            raw_headers=["Account ID", "Order ID", "Amount"],
            records=[
                {"Account ID": "ACME", "Order ID": 1, "Amount": 10},
                {"Account ID": "ACME", "Order ID": 2, "Amount": 20},
            ],
            used_tables=set(),
        )

        self.assertEqual(
            {
                "source_columns": ["Account ID", "Order ID"],
                "columns": ["account_id", "order_number"],
                "format": "ORD-{Account ID}-{Order ID}",
            },
            table["row_key"],
        )

    def test_row_key_format_rejects_rendered_collisions(self):
        config = default_sync_config(slug="orders", file_id="sheet-1")
        config["schema"]["tables"] = {
            "Orders": {
                "row_key": {
                    "columns": ["Account ID", "Order ID"],
                    "format": "{Account ID}-{Order ID}",
                }
            }
        }
        config = validate_sync_config(config, expected_slug="orders")

        with self.assertRaisesRegex(ValueError, "render to the same identifier"):
            loader._record_table(
                config,
                source_name="Orders",
                raw_headers=["Account ID", "Order ID"],
                records=[
                    {"Account ID": "ACME-NORTH", "Order ID": "42"},
                    {"Account ID": "ACME", "Order ID": "NORTH-42"},
                ],
                used_tables=set(),
            )

    def test_row_key_rejects_duplicate_blank_and_disabled_components(self):
        cases = (
            (
                ["Account ID", "Order ID"],
                [
                    {"Account ID": "ACME", "Order ID": 1},
                    {"Account ID": "ACME", "Order ID": 1.0},
                ],
                "not unique",
                {},
            ),
            (
                ["Account ID"],
                [{"Account ID": "   ", "Order ID": 1}],
                "is blank",
                {},
            ),
            (
                ["Account ID"],
                [{"Account ID": "ACME", "Order ID": 1}],
                "missing or disabled",
                {"Account ID": {"enabled": False}},
            ),
        )

        for key_columns, records, message, column_rules in cases:
            with self.subTest(message=message):
                config = default_sync_config(slug="orders", file_id="sheet-1")
                config["schema"]["tables"] = {
                    "Orders": {
                        "row_key": {"columns": key_columns},
                        "columns": column_rules,
                    }
                }
                config = validate_sync_config(config, expected_slug="orders")

                with self.assertRaisesRegex(ValueError, message):
                    loader._record_table(
                        config,
                        source_name="Orders",
                        raw_headers=["Account ID", "Order ID"],
                        records=records,
                        used_tables=set(),
                    )

    def test_google_schema_discovery_returns_headers_without_data_rows(self):
        service = MagicMock()
        service.spreadsheets.return_value.values.return_value.batchGet.return_value.execute.return_value = {
            "valueRanges": [
                {
                    "values": [
                        ["Report"],
                        ["Account ID", "Order ID", "Amount"],
                        ["ACME", 42, 10.5],
                    ]
                },
                {"values": []},
            ]
        }

        with patch("googleapiclient.discovery.build", return_value=service):
            schemas = loader._google_sheet_schemas(
                object(),
                "sheet-1",
                ["Orders", "Empty"],
            )

        self.assertEqual(
            {
                "name": "Orders",
                "header_row": 2,
                "columns": ["Account ID", "Order ID", "Amount"],
                "columns_truncated": False,
            },
            schemas[0],
        )
        self.assertEqual("Empty", schemas[1]["name"])
        self.assertEqual([], schemas[1]["columns"])
        self.assertNotIn("ACME", str(schemas))

    def test_excel_parser_reads_workbook_worksheets(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        workbook.active.title = "Orders"
        workbook.active.append(["Order ID", "Amount"])
        workbook.active.append([1, 10.5])
        workbook.create_sheet("Forecast").append(["Month", "Target"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        worksheets = loader._read_excel_worksheets(
            GoogleDriveFile(
                "excel-1",
                "sales.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            output.getvalue(),
        )

        self.assertEqual(["Orders", "Forecast"], [name for name, _ in worksheets])
        self.assertEqual([1, 10.5], worksheets[0][1][1])
        self.assertEqual(
            ["Orders", "Forecast"],
            loader._read_excel_worksheet_names(
                GoogleDriveFile(
                    "excel-1",
                    "sales.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
                output.getvalue(),
            ),
        )

    def test_excel_discovery_returns_bounded_worksheet_headers(self):
        def configure(workbook):
            orders = workbook.active
            orders.title = "Orders"
            orders.append(["Report"])
            orders.append(["Order ID", "Amount"])
            orders.append([1001, 10.5])
            forecast = workbook.create_sheet("Forecast")
            forecast.append(["Month", "Target"])
            forecast.append(["September", 25])

        discovery = self._discover_file_schema(
            self._excel_file(),
            "excel",
            self._workbook_bytes(configure),
        )

        self.assertEqual(["Orders", "Forecast"], discovery["worksheets"])
        self.assertEqual(
            [
                {
                    "name": "Orders",
                    "header_row": 2,
                    "columns": ["Order ID", "Amount"],
                    "columns_truncated": False,
                },
                {
                    "name": "Forecast",
                    "header_row": 1,
                    "columns": ["Month", "Target"],
                    "columns_truncated": False,
                },
            ],
            discovery["worksheet_schemas"],
        )

    def test_excel_auto_detects_header_after_merged_intro_at_row_fifty(self):
        def configure(workbook):
            sheet = workbook.active
            sheet.title = "Orders"
            sheet.merge_cells("A1:C1")
            sheet["A1"] = "Executive summary"
            sheet.append([])
            sheet.cell(row=50, column=1, value="Order ID")
            sheet.cell(row=50, column=2, value="Customer")
            sheet.cell(row=50, column=3, value="Amount")
            sheet.append([1001, "Ada", 10.25])
            sheet.append([1002, "Grace", 12.5])

        tables, inspection = self._extract_excel(
            self._workbook_bytes(configure),
            self._excel_config(),
        )

        self.assertEqual(50, inspection.header_rows["Orders"])
        self.assertEqual(
            ["order_id", "customer", "amount"],
            [column["name"] for column in tables[0]["columns"]],
        )
        self.assertEqual(
            {"order_id": 1001, "customer": "Ada", "amount": 10.25},
            tables[0]["rows"][0],
        )

    def test_excel_per_worksheet_header_row_can_target_row_fifty(self):
        def configure(workbook):
            sheet = workbook.active
            sheet.title = "Orders"
            sheet["A1"] = "Executive summary"
            sheet.cell(row=50, column=1, value="Order ID")
            sheet.cell(row=50, column=2, value="Amount")
            sheet.append([1001, 10.25])

        tables, inspection = self._extract_excel(
            self._workbook_bytes(configure),
            self._excel_config(
                header_row=1,
                table_rules={"Orders": {"header_row": 50}},
            ),
        )

        self.assertEqual(50, inspection.header_rows["Orders"])
        self.assertEqual(
            {"order_id": 1001, "amount": 10.25},
            tables[0]["rows"][0],
        )

    def test_excel_incorrect_explicit_header_is_not_auto_corrected(self):
        def configure(workbook):
            sheet = workbook.active
            sheet.title = "Orders"
            sheet["A1"] = "Executive summary"
            sheet.cell(row=50, column=1, value="Order ID")
            sheet.cell(row=50, column=2, value="Amount")
            sheet.append([1001, 10.25])

        tables, inspection = self._extract_excel(
            self._workbook_bytes(configure),
            self._excel_config(header_row=1),
        )

        self.assertEqual(1, inspection.header_rows["Orders"])
        self.assertEqual(
            ["executive_summary"],
            [column["name"] for column in tables[0]["columns"]],
        )
        self.assertEqual(
            [{"executive_summary": "Order ID"}, {"executive_summary": 1001}],
            tables[0]["rows"],
        )

    def test_excel_rejects_header_row_beyond_worksheet_bounds(self):
        def configure(workbook):
            sheet = workbook.active
            sheet.title = "Orders"
            sheet.append(["Order ID", "Amount"])
            sheet.append([1001, 10.25])

        with self.assertRaisesRegex(
            ValueError,
            "Configured header row 150 is outside 'Orders'",
        ):
            self._extract_excel(
                self._workbook_bytes(configure),
                self._excel_config(header_row=150),
            )

    def test_excel_skips_empty_worksheets(self):
        def configure(workbook):
            workbook.active.title = "Orders"

        tables, inspection = self._extract_excel(
            self._workbook_bytes(configure),
            self._excel_config(),
        )

        self.assertEqual([], tables)
        self.assertEqual({}, inspection.header_rows)

    def test_excel_normalizes_duplicate_and_blank_headers(self):
        def configure(workbook):
            sheet = workbook.active
            sheet.title = "Orders"
            sheet.append(["Name", "Name", None, "Amount"])
            sheet.append(["Ada", "Lovelace", "ignored", 10.25])

        tables, _inspection = self._extract_excel(
            self._workbook_bytes(configure),
            self._excel_config(header_row=1),
        )

        self.assertEqual(
            ["name", "name_2", "amount"],
            [column["name"] for column in tables[0]["columns"]],
        )
        self.assertEqual(
            {"name": "Ada", "name_2": "Lovelace", "amount": 10.25},
            tables[0]["rows"][0],
        )

    def test_excel_keeps_hidden_worksheet_available_when_selected(self):
        def configure(workbook):
            orders = workbook.active
            orders.title = "Orders"
            orders.append(["Order ID"])
            orders.append([1001])
            archive = workbook.create_sheet("Archive")
            archive.sheet_state = "hidden"
            archive.append(["Order ID"])
            archive.append([9001])

        content = self._workbook_bytes(configure)
        tables, _inspection = self._extract_excel(
            content,
            self._excel_config(sheets=["Orders", "Archive"], header_row=1),
        )

        self.assertEqual(
            ["Orders", "Archive"],
            [table["source_name"] for table in tables],
        )
        self.assertEqual(
            ["Orders", "Archive"],
            loader._read_excel_worksheet_names(self._excel_file(), content),
        )

    def test_excel_formula_without_cached_value_loads_as_null(self):
        def configure(workbook):
            sheet = workbook.active
            sheet.title = "Orders"
            sheet.append(["Order ID", "Calculated Amount"])
            sheet.append([1001, "=10+5"])

        tables, _inspection = self._extract_excel(
            self._workbook_bytes(configure),
            self._excel_config(header_row=1),
        )

        self.assertEqual(
            {"order_id": 1001, "calculated_amount": None},
            tables[0]["rows"][0],
        )

    def test_parquet_parser_preserves_compatible_column_types(self):
        import pyarrow as arrow
        import pyarrow.parquet as parquet

        output = io.BytesIO()
        parquet.write_table(
            arrow.table(
                {
                    "order_id": arrow.array([1, 2], type=arrow.int64()),
                    "amount": arrow.array([10.5, 20.0], type=arrow.float64()),
                    "active": arrow.array([True, False], type=arrow.bool_()),
                }
            ),
            output,
        )
        config = validate_sync_config(
            default_sync_config(slug="orders", file_id="parquet-1"),
            expected_slug="orders",
        )

        tables = loader._extract_parquet(
            config,
            GoogleDriveFile(
                "parquet-1",
                "orders.parquet",
                "application/vnd.apache.parquet",
            ),
            output.getvalue(),
            set(),
        )

        types = {
            column["source_name"]: column["data_type"]
            for column in tables[0]["columns"]
        }
        self.assertEqual(
            {"order_id": "bigint", "amount": "double", "active": "bool"},
            types,
        )
        self.assertEqual(2, len(tables[0]["rows"]))

    def test_parquet_discovery_returns_row_identity_columns(self):
        import pyarrow as arrow
        import pyarrow.parquet as parquet

        output = io.BytesIO()
        parquet.write_table(
            arrow.table(
                {
                    "order_id": arrow.array([1, 2], type=arrow.int64()),
                    "active": arrow.array([True, False], type=arrow.bool_()),
                }
            ),
            output,
        )
        discovery = self._discover_file_schema(
            GoogleDriveFile(
                "parquet-1",
                "orders.parquet",
                "application/vnd.apache.parquet",
            ),
            "parquet",
            output.getvalue(),
        )

        self.assertEqual([], discovery["worksheets"])
        self.assertEqual(
            [
                {
                    "name": "orders",
                    "columns": ["order_id", "active"],
                    "columns_truncated": False,
                }
            ],
            discovery["worksheet_schemas"],
        )


if __name__ == "__main__":
    unittest.main()
