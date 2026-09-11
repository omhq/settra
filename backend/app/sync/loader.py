from __future__ import annotations

import asyncio
import fnmatch
import io
import json
import logging
import os
import re
import unicodedata

from datetime import date, datetime, timezone
from decimal import Decimal
from numbers import Number
from pathlib import PurePath
from typing import Any

import aiofiles
import yaml

from fastapi import HTTPException

from app.auth import current_organization_id
from app.db import db_connection
from app.destinations import DestinationRuntime, runtime_from_connection
from app.routers.constants import (
    DLT_PIPELINES_DIR,
    GOOGLE_DRIVE_KEY,
)
from app.sync.config import (
    MAX_RENDERED_ROW_KEY_LENGTH,
    apply_detected_source_config,
    config_path,
    read_sync_config,
    render_row_key_format,
    row_key_columns,
    row_key_format,
    table_rule,
    write_sync_config,
)
from app.sync.inspection import (
    GoogleDriveFile,
    TabularFileInspector,
    TabularInspection,
    display_delimiter,
    resolve_header_row,
)
from app.sync.secrets import load_google_oauth_secret

logger = logging.getLogger(__name__)

GOOGLE_FILE_SCOPE = "https://www.googleapis.com/auth/drive.file"
SCHEMA_DISCOVERY_MAX_TABLES = 20
SCHEMA_DISCOVERY_MAX_ROWS = 100
SCHEMA_DISCOVERY_MAX_COLUMNS = 100
_SYNC_LOCKS: dict[int, asyncio.Lock] = {}


async def run_connection_sync(
    connection_id: int,
    *,
    trigger: str = "manual",
    organization_id: int | None = None,
) -> dict[str, Any]:
    lock = _SYNC_LOCKS.setdefault(connection_id, asyncio.Lock())

    if lock.locked():
        raise HTTPException(409, "A sync is already running for this connection")

    async with lock:
        effective_organization_id = (
            organization_id
            if organization_id is not None
            else current_organization_id()
        )
        connection = await _connection(connection_id, effective_organization_id)
        storage_key = connection.get("storage_key") or connection["slug"]
        config = await read_sync_config(
            storage_key,
            expected_destination_key=(
                connection.get("destination_slug") or "built_in_postgres"
            ),
            expected_destination_schema=(
                connection.get("destination_schema") or storage_key
            ),
        )

        if not config:
            raise HTTPException(409, "Connection sync configuration is missing")

        secret = await load_google_oauth_secret(effective_organization_id)

        if GOOGLE_FILE_SCOPE not in set(secret.get("scopes") or []):
            raise HTTPException(
                409,
                "Reconnect Google to grant file-specific Drive access",
            )

        run_id = await _start_run(connection_id, trigger)
        result: dict[str, Any] | None = None

        try:
            result = await asyncio.to_thread(
                _run_dlt_sync,
                connection,
                config,
                secret,
            )
            detected_source = result.pop("_detected_source", None)

            if isinstance(detected_source, dict):
                config = apply_detected_source_config(config, detected_source)
                await write_sync_config(
                    storage_key,
                    config,
                    expected_destination_key=(
                        connection.get("destination_slug") or "built_in_postgres"
                    ),
                    expected_destination_schema=(
                        connection.get("destination_schema") or storage_key
                    ),
                )

            await _refresh_models_and_metadata(connection)
            await _finish_run(run_id, connection_id, result=result)
            return {"ok": True, "run_id": run_id, **result}
        except HTTPException as exc:
            await _finish_run(
                run_id,
                connection_id,
                result=result,
                error=str(exc.detail),
            )
            raise
        except Exception as exc:
            logger.exception(
                "Google Drive sync failed connection_id=%s trigger=%s",
                connection_id,
                trigger,
            )
            message = f"{exc.__class__.__name__}: {exc}"
            await _finish_run(
                run_id,
                connection_id,
                result=result,
                error=message,
            )
            raise HTTPException(502, f"Drive file sync failed: {message}") from exc


def _run_dlt_sync(
    connection: dict[str, Any],
    config: dict[str, Any],
    secret: dict[str, Any],
) -> dict[str, Any]:
    # Imports stay local so metadata-only commands do not need the loader runtime.
    import dlt

    from dlt.sources.credentials import GcpOAuthCredentials

    os.environ.setdefault("RUNTIME__DLTHUB_TELEMETRY", "false")
    DLT_PIPELINES_DIR.mkdir(parents=True, exist_ok=True)

    oauth = GcpOAuthCredentials(
        client_id=_google_client_id(),
        client_secret=_google_client_secret(),
        refresh_token=str(secret["refresh_token"]),
        project_id=os.getenv("GOOGLE_CLOUD_PROJECT") or None,
        scopes=[GOOGLE_FILE_SCOPE],
    )
    destination_runtime = runtime_from_connection(connection)
    destination_runtime.require_built_in_postgres()
    try:
        extracted, inspection = _extract_google_drive_file(config, oauth)
    except Exception as exc:
        from googleapiclient.errors import HttpError

        if isinstance(exc, HttpError) and getattr(exc.resp, "status", None) == 403:
            raise HTTPException(
                403,
                "The selected Drive file is not authorized for Settra. "
                "Choose it again through Google Picker.",
            ) from exc

        raise

    if not extracted:
        raise ValueError("No selected tables contain a usable header row or schema")

    resources = []
    loaded_tables: list[dict[str, Any]] = []

    for table in extracted:
        column_hints = {
            column["name"]: {
                **(
                    {"data_type": column.get("data_type") or "text"}
                    if not table["rows"] or column.get("data_type")
                    else {}
                ),
                "nullable": bool(column.get("nullable", True)),
            }
            for column in table["columns"]
        }
        resource = dlt.resource(
            table["rows"],
            name=table["table_name"],
            write_disposition="replace",
            columns=column_hints,
            schema_contract=config["load"].get("schema_contract"),
        )
        resources.append(resource)
        loaded_tables.append(
            {key: value for key, value in table.items() if key != "rows"}
            | {"row_count": len(table["rows"])}
        )

    @dlt.source(name=f"settra_{connection['storage_key']}_google_drive")
    def source():
        return resources

    destination = dlt.destinations.postgres(
        credentials=destination_runtime.postgres_dsn(),
        replace_strategy="insert-from-staging",
    )
    pipeline = dlt.pipeline(
        pipeline_name=f"settra_{connection['storage_key']}",
        destination=destination,
        dataset_name=destination_runtime.schema,
        pipelines_dir=str(DLT_PIPELINES_DIR),
    )
    load_info = pipeline.run(source())
    load_ids = list(getattr(load_info, "loads_ids", []) or [])

    _finalize_postgres_schema(
        destination=destination_runtime,
        schema=destination_runtime.schema,
        manifest_slug=connection["storage_key"],
        tables=loaded_tables,
        config=config,
    )
    manifest = _postgres_manifest(
        connection=connection,
        loaded_tables=loaded_tables,
        load_ids=load_ids,
        source_info=inspection.manifest_source(),
        destination=destination_runtime,
    )
    _write_manifest(connection["storage_key"], manifest)

    return {
        "connection_id": connection["id"],
        "schema": destination_runtime.schema,
        "destination_id": destination_runtime.id,
        "tables": manifest["tables"],
        "table_count": len(manifest["tables"]),
        "row_count": sum(item["row_count"] for item in loaded_tables),
        "load_ids": load_ids,
        "source_format": inspection.format,
        "completed_at": manifest["generated_at"],
        "_detected_source": {
            "file_name": inspection.file.name,
            "mime_type": inspection.file.mime_type,
            "format": inspection.format,
            "parsing": inspection.parsing,
            "header_rows": inspection.header_rows,
        },
    }


def _extract_google_drive_file(
    config: dict[str, Any],
    oauth: Any,
) -> tuple[list[dict[str, Any]], TabularInspection]:
    from googleapiclient.discovery import build

    native_credentials = oauth.to_native_credentials()
    drive = build(
        "drive",
        "v3",
        credentials=native_credentials,
        cache_discovery=False,
    )
    source = config["source"]
    file, source_format, content = _prepare_google_drive_file(
        drive,
        source["file_id"],
        configured_format=str(source.get("format") or "auto"),
    )
    inspector = TabularFileInspector()

    inspection = TabularInspection(file=file, format=source_format)
    used_tables: set[str] = set()

    if source_format == "google_sheets":
        tables = _extract_google_sheets(
            config,
            native_credentials,
            inspector,
            inspection,
            used_tables,
        )
    elif source_format == "csv":
        tables = _extract_csv(
            config,
            file,
            content or b"",
            inspector,
            inspection,
            used_tables,
        )
    elif source_format == "excel":
        tables = _extract_excel(
            config,
            file,
            content or b"",
            inspector,
            inspection,
            used_tables,
        )
    elif source_format == "parquet":
        tables = _extract_parquet(config, file, content or b"", used_tables)
    else:  # pragma: no cover - config validation and inspector guard this.
        raise ValueError(f"Unsupported source format: {source_format}")

    return tables, inspection


def discover_google_drive_worksheets(
    file_id: str,
    native_credentials: Any,
) -> dict[str, Any]:
    """Return table names and bounded header metadata, never data rows."""

    from googleapiclient.discovery import build

    drive = build(
        "drive",
        "v3",
        credentials=native_credentials,
        cache_discovery=False,
    )
    file, source_format, content = _prepare_google_drive_file(
        drive,
        file_id,
        content_formats={"csv", "excel", "parquet"},
    )

    if source_format == "google_sheets":
        worksheets = _google_sheet_titles(native_credentials, file.file_id)
        worksheet_schemas = _google_sheet_schemas(
            native_credentials,
            file.file_id,
            worksheets[:SCHEMA_DISCOVERY_MAX_TABLES],
        )
    elif source_format == "excel":
        worksheets, worksheet_schemas = _excel_worksheet_schemas(
            file,
            content or b"",
        )
    elif source_format == "csv":
        worksheets = []
        worksheet_schemas = [
            _csv_file_schema(file, content or b"", TabularFileInspector())
        ]
    elif source_format == "parquet":
        worksheets = []
        worksheet_schemas = [_parquet_file_schema(file, content or b"")]
    else:
        worksheets = []
        worksheet_schemas = []

    return {
        "file_name": file.name,
        "mime_type": file.mime_type,
        "format": source_format,
        "worksheets": worksheets,
        "worksheet_schemas": worksheet_schemas,
        "worksheet_schema_truncated": (
            source_format in {"google_sheets", "excel"}
            and len(worksheets) > SCHEMA_DISCOVERY_MAX_TABLES
        ),
    }


def _prepare_google_drive_file(
    drive: Any,
    file_id: str,
    *,
    configured_format: str = "auto",
    content_formats: set[str] | None = None,
) -> tuple[GoogleDriveFile, str, bytes | None]:
    file = _drive_file_metadata(drive, file_id)
    inspector = TabularFileInspector()
    content: bytes | None = None

    try:
        source_format = inspector.detect_format(file, configured=configured_format)
    except ValueError:
        if not file.can_download:
            raise

        content = _download_drive_file(drive, file.file_id)
        source_format = inspector.detect_format(
            file,
            configured=configured_format,
            content=content,
        )

    required_content_formats = content_formats or {"csv", "excel", "parquet"}

    if source_format in required_content_formats:
        if not file.can_download:
            raise ValueError("The selected Drive file does not allow downloads")

        if content is None:
            content = _download_drive_file(drive, file.file_id)
        source_format = inspector.detect_format(
            file,
            configured=configured_format,
            content=content,
        )

    return file, source_format, content


def _drive_file_metadata(drive: Any, file_id: str) -> GoogleDriveFile:
    item = (
        drive.files()
        .get(
            fileId=file_id,
            supportsAllDrives=True,
            fields=(
                "id,name,mimeType,size,modifiedTime,md5Checksum,trashed,"
                "capabilities(canDownload),shortcutDetails(targetId,targetMimeType)"
            ),
        )
        .execute()
    )

    if item.get("trashed"):
        raise ValueError("The selected Drive file is in the trash")
    if item.get("mimeType") == "application/vnd.google-apps.shortcut":
        raise ValueError(
            "Google Drive shortcuts are not supported; select the target file directly"
        )

    raw_size = item.get("size")
    return GoogleDriveFile(
        file_id=str(item.get("id") or file_id),
        name=str(item.get("name") or "drive_file"),
        mime_type=str(item.get("mimeType") or "application/octet-stream"),
        size=int(raw_size) if raw_size is not None else None,
        modified_time=str(item.get("modifiedTime") or "") or None,
        checksum=str(item.get("md5Checksum") or "") or None,
        can_download=bool(item.get("capabilities", {}).get("canDownload", True)),
    )


def _download_drive_file(drive: Any, file_id: str) -> bytes:
    from googleapiclient.http import MediaIoBaseDownload

    output = io.BytesIO()
    request = drive.files().get_media(fileId=file_id, supportsAllDrives=True)
    downloader = MediaIoBaseDownload(output, request, chunksize=8 * 1024 * 1024)
    done = False

    while not done:
        _, done = downloader.next_chunk()

    return output.getvalue()


def _extract_google_sheets(
    config: dict[str, Any],
    native_credentials: Any,
    inspector: TabularFileInspector,
    inspection: TabularInspection,
    used_tables: set[str],
) -> list[dict[str, Any]]:
    source = config["source"]
    file_id = source["file_id"]
    patterns = source.get("sheets") or ["*"]
    selected = [
        title
        for title in _google_sheet_titles(native_credentials, file_id)
        if _matches_any(title, patterns)
    ]

    if not selected:
        raise ValueError("The configured sheet patterns matched no Google Sheet tabs")

    extracted: list[dict[str, Any]] = []

    from googleapiclient.discovery import build

    service = build(
        "sheets",
        "v4",
        credentials=native_credentials,
        cache_discovery=False,
    )

    for sheet_title in selected:
        rule = table_rule(config, sheet_title)

        if rule.get("enabled") is False:
            continue

        quoted_range = _quoted_sheet_title(sheet_title)
        response = (
            service.spreadsheets()
            .values()
            .get(
                spreadsheetId=file_id,
                range=quoted_range,
                majorDimension="ROWS",
                valueRenderOption="UNFORMATTED_VALUE",
                dateTimeRenderOption="FORMATTED_STRING",
            )
            .execute()
        )
        values = response.get("values") or []

        if not values:
            continue

        header_row = _resolved_table_header_row(
            config,
            sheet_title,
            values,
            inspector,
        )
        inspection.header_rows[sheet_title] = header_row
        table = _matrix_table(
            config,
            source_name=sheet_title,
            values=values,
            header_row=header_row,
            used_tables=used_tables,
        )

        if table:
            extracted.append(table)

    return extracted


def _google_sheet_titles(native_credentials: Any, file_id: str) -> list[str]:
    from googleapiclient.discovery import build

    service = build(
        "sheets",
        "v4",
        credentials=native_credentials,
        cache_discovery=False,
    )
    workbook = (
        service.spreadsheets()
        .get(
            spreadsheetId=file_id,
            fields="sheets(properties(title,index,sheetType))",
        )
        .execute()
    )
    return [
        item["properties"]["title"]
        for item in workbook.get("sheets", [])
        if item.get("properties", {}).get("sheetType", "GRID") == "GRID"
        and item.get("properties", {}).get("title")
    ]


def _google_sheet_schemas(
    native_credentials: Any,
    file_id: str,
    sheet_titles: list[str],
) -> list[dict[str, Any]]:
    """Discover bounded header metadata without returning worksheet values."""

    if not sheet_titles:
        return []

    from googleapiclient.discovery import build

    service = build(
        "sheets",
        "v4",
        credentials=native_credentials,
        cache_discovery=False,
    )
    ranges = [
        f"{_quoted_sheet_title(title)}!A1:CV{SCHEMA_DISCOVERY_MAX_ROWS}"
        for title in sheet_titles
    ]
    response = (
        service.spreadsheets()
        .values()
        .batchGet(
            spreadsheetId=file_id,
            ranges=ranges,
            majorDimension="ROWS",
            valueRenderOption="UNFORMATTED_VALUE",
            dateTimeRenderOption="FORMATTED_STRING",
        )
        .execute()
    )
    value_ranges = response.get("valueRanges") or []
    inspector = TabularFileInspector()
    schemas: list[dict[str, Any]] = []

    for index, title in enumerate(sheet_titles):
        value_range = value_ranges[index] if index < len(value_ranges) else {}
        values = value_range.get("values") or []
        schemas.append(
            _table_schema_from_rows(
                title,
                values,
                inspector,
                columns_truncated_at_limit=True,
            )
        )

    return schemas


def _table_schema_from_rows(
    name: str,
    rows: list[list[Any]],
    inspector: TabularFileInspector,
    *,
    header_row: int | None = None,
    columns_truncated: bool = False,
    columns_truncated_at_limit: bool = False,
) -> dict[str, Any]:
    try:
        resolved_header_row = header_row or inspector.detect_header_row(rows)
        header_values = rows[resolved_header_row - 1]
        columns = [
            header["source_name"]
            for header in _headers(header_values[:SCHEMA_DISCOVERY_MAX_COLUMNS])
        ]

        if not columns:
            raise ValueError("No usable header columns were found")

        return {
            "name": name,
            "header_row": resolved_header_row,
            "columns": columns,
            "columns_truncated": columns_truncated
            or len(header_values) > SCHEMA_DISCOVERY_MAX_COLUMNS
            or (
                columns_truncated_at_limit
                and len(header_values) >= SCHEMA_DISCOVERY_MAX_COLUMNS
            ),
        }
    except ValueError as exc:
        return {
            "name": name,
            "columns": [],
            "error": str(exc),
        }


def _quoted_sheet_title(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _extract_csv(
    config: dict[str, Any],
    file: GoogleDriveFile,
    content: bytes,
    inspector: TabularFileInspector,
    inspection: TabularInspection,
    used_tables: set[str],
) -> list[dict[str, Any]]:
    parsing = config["source"].get("parsing") or {}
    encoding, delimiter, values, detected_header_row = inspector.inspect_delimited(
        content,
        file_name=file.name,
        parsing=parsing,
    )
    source_name = _file_table_name(file.name)
    rule = table_rule(config, source_name)
    header_row = resolve_header_row(
        rule.get("header_row", parsing.get("header_row", detected_header_row)),
        values,
        inspector,
    )
    inspection.parsing = {
        "encoding": encoding,
        "delimiter": display_delimiter(delimiter),
        "header_row": header_row,
    }
    inspection.header_rows[source_name] = header_row
    table = _matrix_table(
        config,
        source_name=source_name,
        values=values,
        header_row=header_row,
        used_tables=used_tables,
    )
    return [table] if table else []


def _csv_file_schema(
    file: GoogleDriveFile,
    content: bytes,
    inspector: TabularFileInspector,
) -> dict[str, Any]:
    source_name = _file_table_name(file.name)

    try:
        _, _, rows, header_row = inspector.inspect_delimited(
            content,
            file_name=file.name,
            parsing={
                "encoding": "auto",
                "delimiter": "auto",
                "header_row": "auto",
            },
        )
    except ValueError as exc:
        return {"name": source_name, "columns": [], "error": str(exc)}

    return _table_schema_from_rows(
        source_name,
        rows,
        inspector,
        header_row=header_row,
    )


def _extract_excel(
    config: dict[str, Any],
    file: GoogleDriveFile,
    content: bytes,
    inspector: TabularFileInspector,
    inspection: TabularInspection,
    used_tables: set[str],
) -> list[dict[str, Any]]:
    worksheets = _read_excel_worksheets(file, content)
    patterns = config["source"].get("sheets") or ["*"]
    selected = [
        (name, rows) for name, rows in worksheets if _matches_any(name, patterns)
    ]

    if not selected:
        raise ValueError("The configured sheet patterns matched no Excel worksheets")

    extracted: list[dict[str, Any]] = []

    for sheet_title, values in selected:
        rule = table_rule(config, sheet_title)

        if rule.get("enabled") is False or not values:
            continue

        header_row = _resolved_table_header_row(
            config,
            sheet_title,
            values,
            inspector,
        )
        inspection.header_rows[sheet_title] = header_row
        table = _matrix_table(
            config,
            source_name=sheet_title,
            values=values,
            header_row=header_row,
            used_tables=used_tables,
        )

        if table:
            extracted.append(table)

    return extracted


def _read_excel_worksheets(
    file: GoogleDriveFile,
    content: bytes,
) -> list[tuple[str, list[list[Any]]]]:
    is_legacy_xls = file.name.lower().endswith(".xls") or content.startswith(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    )

    if is_legacy_xls:
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - installed in runtime image.
            raise ValueError("Legacy Excel support requires the xlrd package") from exc

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        worksheets: list[tuple[str, list[list[Any]]]] = []

        try:
            for sheet in workbook.sheets():
                rows = []

                for row_index in range(sheet.nrows):
                    values = []

                    for cell in sheet.row(row_index):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            values.append(
                                xlrd.xldate.xldate_as_datetime(
                                    cell.value,
                                    workbook.datemode,
                                )
                            )
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            values.append(bool(cell.value))
                        elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                            values.append(None)
                        else:
                            values.append(cell.value)

                    rows.append(values)

                worksheets.append((sheet.name, rows))
        finally:
            workbook.release_resources()

        return worksheets

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - installed in runtime image.
        raise ValueError("Excel support requires the openpyxl package") from exc

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(f"Excel workbook could not be opened: {exc}") from exc

    try:
        return [
            (
                worksheet.title,
                [list(row) for row in worksheet.iter_rows(values_only=True)],
            )
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _excel_worksheet_schemas(
    file: GoogleDriveFile,
    content: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    """Inspect bounded worksheet headers without returning workbook rows."""

    is_legacy_xls = file.name.lower().endswith(".xls") or content.startswith(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    )
    inspector = TabularFileInspector()

    if is_legacy_xls:
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - installed in runtime image.
            raise ValueError("Legacy Excel support requires the xlrd package") from exc

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)

        try:
            worksheets = list(workbook.sheet_names())
            schemas = []

            for name in worksheets[:SCHEMA_DISCOVERY_MAX_TABLES]:
                sheet = workbook.sheet_by_name(name)
                rows = [
                    [
                        sheet.cell_value(row_index, column_index)
                        for column_index in range(
                            min(sheet.ncols, SCHEMA_DISCOVERY_MAX_COLUMNS)
                        )
                    ]
                    for row_index in range(min(sheet.nrows, SCHEMA_DISCOVERY_MAX_ROWS))
                ]
                schemas.append(
                    _table_schema_from_rows(
                        str(name),
                        rows,
                        inspector,
                        columns_truncated=sheet.ncols > SCHEMA_DISCOVERY_MAX_COLUMNS,
                    )
                )

            return [str(name) for name in worksheets], schemas
        finally:
            workbook.release_resources()

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - installed in runtime image.
        raise ValueError("Excel support requires the openpyxl package") from exc

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(f"Excel workbook could not be opened: {exc}") from exc

    try:
        worksheets = list(workbook.sheetnames)
        schemas = []

        for name in worksheets[:SCHEMA_DISCOVERY_MAX_TABLES]:
            worksheet = workbook[name]
            max_row = min(worksheet.max_row or 1, SCHEMA_DISCOVERY_MAX_ROWS)
            max_column = min(
                worksheet.max_column or 1,
                SCHEMA_DISCOVERY_MAX_COLUMNS,
            )
            rows = [
                list(row)
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    max_col=max_column,
                    values_only=True,
                )
            ]
            schemas.append(
                _table_schema_from_rows(
                    name,
                    rows,
                    inspector,
                    columns_truncated=(
                        worksheet.max_column > SCHEMA_DISCOVERY_MAX_COLUMNS
                    ),
                )
            )

        return worksheets, schemas
    finally:
        workbook.close()


def _read_excel_worksheet_names(
    file: GoogleDriveFile,
    content: bytes,
) -> list[str]:
    is_legacy_xls = file.name.lower().endswith(".xls") or content.startswith(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    )

    if is_legacy_xls:
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - installed in runtime image.
            raise ValueError("Legacy Excel support requires the xlrd package") from exc

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)

        try:
            return [str(name) for name in workbook.sheet_names()]
        finally:
            workbook.release_resources()

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - installed in runtime image.
        raise ValueError("Excel support requires the openpyxl package") from exc

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(f"Excel workbook could not be opened: {exc}") from exc

    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _extract_parquet(
    config: dict[str, Any],
    file: GoogleDriveFile,
    content: bytes,
    used_tables: set[str],
) -> list[dict[str, Any]]:
    try:
        import pyarrow.parquet as parquet
    except ImportError as exc:  # pragma: no cover - installed in runtime image.
        raise ValueError("Parquet support requires the pyarrow package") from exc

    try:
        parquet_file = parquet.ParquetFile(io.BytesIO(content))
        # Arrow's threaded Parquet reader can race with Python interpreter
        # shutdown and abort an otherwise successful short-lived process.
        arrow_table = parquet_file.read(use_threads=False)
    except Exception as exc:
        raise ValueError(f"Parquet file could not be read: {exc}") from exc

    source_name = _file_table_name(file.name)
    inferred_types = {
        field.name: _arrow_dlt_type(field.type) for field in arrow_table.schema
    }
    table = _record_table(
        config,
        source_name=source_name,
        raw_headers=arrow_table.column_names,
        records=arrow_table.to_pylist(),
        used_tables=used_tables,
        inferred_types=inferred_types,
    )
    return [table] if table else []


def _parquet_file_schema(
    file: GoogleDriveFile,
    content: bytes,
) -> dict[str, Any]:
    source_name = _file_table_name(file.name)

    try:
        import pyarrow.parquet as parquet
    except ImportError as exc:  # pragma: no cover - installed in runtime image.
        raise ValueError("Parquet support requires the pyarrow package") from exc

    try:
        fields = list(parquet.ParquetFile(io.BytesIO(content)).schema_arrow)
    except Exception as exc:
        return {
            "name": source_name,
            "columns": [],
            "error": f"Parquet file could not be opened: {exc}",
        }

    columns = [
        header["source_name"]
        for header in _headers(
            [field.name for field in fields[:SCHEMA_DISCOVERY_MAX_COLUMNS]]
        )
    ]
    return {
        "name": source_name,
        "columns": columns,
        "columns_truncated": len(fields) > SCHEMA_DISCOVERY_MAX_COLUMNS,
    }


def _arrow_dlt_type(value: Any) -> str:
    import pyarrow.types as arrow_types

    if arrow_types.is_boolean(value):
        return "bool"
    if arrow_types.is_binary(value) or arrow_types.is_large_binary(value):
        return "binary"
    if arrow_types.is_integer(value):
        return "bigint"
    if arrow_types.is_floating(value):
        return "double"
    if arrow_types.is_decimal(value):
        return "decimal"
    if arrow_types.is_date(value):
        return "date"
    if arrow_types.is_timestamp(value) or arrow_types.is_time(value):
        return "timestamp"
    if (
        arrow_types.is_list(value)
        or arrow_types.is_large_list(value)
        or arrow_types.is_fixed_size_list(value)
        or arrow_types.is_struct(value)
        or arrow_types.is_map(value)
    ):
        return "json"

    return "text"


def _resolved_table_header_row(
    config: dict[str, Any],
    source_name: str,
    values: list[list[Any]],
    inspector: TabularFileInspector,
) -> int:
    parsing = config["source"].get("parsing") or {}
    rule = table_rule(config, source_name)
    return resolve_header_row(
        rule.get("header_row", parsing.get("header_row", "auto")),
        values,
        inspector,
    )


def _matrix_table(
    config: dict[str, Any],
    *,
    source_name: str,
    values: list[list[Any]],
    header_row: int,
    used_tables: set[str],
) -> dict[str, Any] | None:
    header_index = header_row - 1

    if header_index >= len(values):
        raise ValueError(
            f"Configured header row {header_row} is outside {source_name!r}"
        )

    headers = _headers(values[header_index])
    records = [
        {
            header["source_name"]: (
                raw_row[header["source_index"]]
                if header["source_index"] < len(raw_row)
                else None
            )
            for header in headers
        }
        for raw_row in values[header_index + 1 :]
        if any(value is not None and str(value).strip() for value in raw_row)
    ]
    return _record_table(
        config,
        source_name=source_name,
        raw_headers=[header["source_name"] for header in headers],
        records=records,
        used_tables=used_tables,
    )


def _record_table(
    config: dict[str, Any],
    *,
    source_name: str,
    raw_headers: list[Any],
    records: list[dict[str, Any]],
    used_tables: set[str],
    inferred_types: dict[str, str] | None = None,
) -> dict[str, Any] | None:
    rule = table_rule(config, source_name)

    if rule.get("enabled") is False:
        return None

    headers = _headers(raw_headers)

    if not headers:
        return None

    table_name = _unique_name(
        _identifier(str(rule.get("table_name") or source_name), "table"),
        used_tables,
    )
    used_tables.add(table_name)
    column_rules = rule.get("columns") if isinstance(rule.get("columns"), dict) else {}
    used_columns: set[str] = set()
    columns: list[dict[str, Any]] = []

    for index, header in enumerate(headers):
        raw_rule = column_rules.get(header["source_name"], {})
        column_rule = raw_rule if isinstance(raw_rule, dict) else {}

        if column_rule.get("enabled") is False:
            continue

        name = _unique_name(
            _identifier(
                str(column_rule.get("name") or header["source_name"]),
                f"column_{index + 1}",
            ),
            used_columns,
        )
        used_columns.add(name)
        detected_type = (inferred_types or {}).get(header["source_name"])
        columns.append(
            {
                "source_name": header["source_name"],
                "name": name,
                "description": str(column_rule.get("description") or ""),
                **(
                    {"data_type": str(column_rule.get("data_type") or detected_type)}
                    if column_rule.get("data_type") or detected_type
                    else {}
                ),
                "nullable": bool(column_rule.get("nullable", True)),
            }
        )

    if not columns:
        return None

    configured_row_key = row_key_columns(rule)
    row_key = _validated_table_row_key(
        source_name,
        configured_row_key,
        row_key_format(rule),
        columns,
        records,
    )

    rows = [
        {column["name"]: record.get(column["source_name"]) for column in columns}
        for record in records
    ]
    return {
        "source_name": source_name,
        "table_name": table_name,
        "description": str(rule.get("description") or ""),
        "columns": columns,
        "rows": rows,
        **({"row_key": row_key} if row_key else {}),
    }


def _validated_table_row_key(
    source_name: str,
    configured_columns: list[str],
    configured_format: str | None,
    columns: list[dict[str, Any]],
    records: list[dict[str, Any]],
) -> dict[str, Any] | None:
    if not configured_columns:
        return None

    destination_by_source = {
        str(column["source_name"]): str(column["name"]) for column in columns
    }
    missing = [
        column for column in configured_columns if column not in destination_by_source
    ]

    if missing:
        raise ValueError(
            f"Configured row key for {source_name!r} references missing or disabled "
            f"columns: {', '.join(missing)}"
        )

    seen: dict[tuple[Any, ...], int] = {}
    formatted_seen: dict[str, int] = {}

    for position, record in enumerate(records, start=1):
        components = []
        formatted_values: dict[str, str] = {}

        for column in configured_columns:
            value = record.get(column)

            if value is None or (isinstance(value, str) and not value.strip()):
                raise ValueError(
                    f"Configured row key for {source_name!r} is blank in data row "
                    f"{position} at column {column!r}"
                )

            components.append(_row_key_component(value))
            formatted_values[column] = _row_key_string_component(value)

        key = tuple(components)
        previous_position = seen.get(key)

        if previous_position is not None:
            raise ValueError(
                f"Configured row key for {source_name!r} is not unique; data rows "
                f"{previous_position} and {position} match"
            )

        seen[key] = position

        if configured_format is not None:
            formatted_key = render_row_key_format(
                configured_format,
                formatted_values,
            )
            if len(formatted_key) > MAX_RENDERED_ROW_KEY_LENGTH:
                raise ValueError(
                    f"Configured row-key format for {source_name!r} exceeds "
                    f"{MAX_RENDERED_ROW_KEY_LENGTH} characters in data row "
                    f"{position}"
                )
            previous_position = formatted_seen.get(formatted_key)

            if previous_position is not None:
                raise ValueError(
                    f"Configured row-key format for {source_name!r} is not unique; "
                    f"data rows {previous_position} and {position} render to the "
                    "same identifier"
                )

            formatted_seen[formatted_key] = position

    return {
        "source_columns": configured_columns,
        "columns": [destination_by_source[column] for column in configured_columns],
        **({"format": configured_format} if configured_format is not None else {}),
    }


def _row_key_component(value: Any) -> tuple[str, Any]:
    if isinstance(value, bool):
        return ("boolean", value)
    if isinstance(value, Number):
        try:
            return ("number", Decimal(str(value)).normalize())
        except Exception:
            return ("number", str(value))
    if isinstance(value, str):
        return ("string", value)
    if isinstance(value, bytes):
        return ("binary", value)

    return (
        value.__class__.__name__,
        json.dumps(value, sort_keys=True, separators=(",", ":"), default=str),
    )


def _row_key_string_component(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Number):
        try:
            return format(Decimal(str(value)).normalize(), "f")
        except Exception:
            return str(value)
    if isinstance(value, str):
        return value
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if not isinstance(value, (dict, list, tuple)):
        return str(value)

    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)


def _file_table_name(file_name: str) -> str:
    stem = PurePath(file_name).stem.strip()
    return stem or "data"


def _headers(raw_headers: list[Any]) -> list[dict[str, Any]]:
    used: set[str] = set()
    headers = []

    for index, value in enumerate(raw_headers):
        source_name = "" if value is None else str(value).strip()

        if not source_name:
            continue

        source_name = _unique_name(source_name, used)
        used.add(source_name)
        headers.append({"source_name": source_name, "source_index": index})

    return headers


def _identifier(value: str, fallback: str) -> str:
    ascii_value = (
        unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    )
    normalized = re.sub(r"[^a-z0-9]+", "_", ascii_value.lower()).strip("_")

    if not normalized:
        normalized = fallback

    if normalized[0].isdigit():
        normalized = f"col_{normalized}"

    return normalized[:63].rstrip("_")


def _unique_name(value: str, used: set[str]) -> str:
    candidate = value
    suffix = 2

    while candidate in used:
        ending = f"_{suffix}"
        candidate = f"{value[: 63 - len(ending)]}{ending}"
        suffix += 1

    return candidate


def _matches_any(title: str, patterns: list[str]) -> bool:
    return any(fnmatch.fnmatchcase(title, pattern) for pattern in patterns)


def _finalize_postgres_schema(
    *,
    destination: DestinationRuntime,
    schema: str,
    manifest_slug: str,
    tables: list[dict[str, Any]],
    config: dict[str, Any],
) -> None:
    import psycopg2

    from psycopg2 import sql

    connection = psycopg2.connect(**destination.postgres_connect_kwargs())

    try:
        with connection.cursor() as cursor:
            previous_tables = _previous_manifest_table_names(manifest_slug)
            current_tables = {item["table_name"] for item in tables}

            for stale in sorted(previous_tables - current_tables):
                cursor.execute(
                    sql.SQL("DROP TABLE IF EXISTS {}.{} CASCADE").format(
                        sql.Identifier(schema),
                        sql.Identifier(stale),
                    )
                )

            for table in tables:
                if not table.get("rows") and not table.get("row_count"):
                    cursor.execute(
                        sql.SQL("CREATE SCHEMA IF NOT EXISTS {}").format(
                            sql.Identifier(schema)
                        )
                    )
                    definitions = []

                    for column in table["columns"]:
                        definition = sql.SQL("{} {}").format(
                            sql.Identifier(column["name"]),
                            sql.SQL(_postgres_type(column.get("data_type") or "text")),
                        )

                        if not column.get("nullable", True):
                            definition += sql.SQL(" NOT NULL")

                        definitions.append(definition)

                    cursor.execute(
                        sql.SQL("CREATE TABLE IF NOT EXISTS {}.{} ({})").format(
                            sql.Identifier(schema),
                            sql.Identifier(table["table_name"]),
                            sql.SQL(", ").join(definitions),
                        )
                    )
                    cursor.execute(
                        sql.SQL("TRUNCATE TABLE {}.{}").format(
                            sql.Identifier(schema),
                            sql.Identifier(table["table_name"]),
                        )
                    )

                cursor.execute(
                    sql.SQL("COMMENT ON TABLE {}.{} IS %s").format(
                        sql.Identifier(schema),
                        sql.Identifier(table["table_name"]),
                    ),
                    (table.get("description") or None,),
                )

                for column in table["columns"]:
                    cursor.execute(
                        sql.SQL("COMMENT ON COLUMN {}.{}.{} IS %s").format(
                            sql.Identifier(schema),
                            sql.Identifier(table["table_name"]),
                            sql.Identifier(column["name"]),
                        ),
                        (column.get("description") or None,),
                    )

        connection.commit()
    finally:
        connection.close()


def _postgres_type(data_type: str) -> str:
    return {
        "binary": "bytea",
        "text": "text",
        "bigint": "bigint",
        "double": "double precision",
        "bool": "boolean",
        "timestamp": "timestamp with time zone",
        "date": "date",
        "decimal": "numeric",
        "json": "jsonb",
    }[data_type]


def _postgres_manifest(
    *,
    connection: dict[str, Any],
    loaded_tables: list[dict[str, Any]],
    load_ids: list[str],
    source_info: dict[str, Any],
    destination: DestinationRuntime,
) -> dict[str, Any]:
    import psycopg2

    connection_pg = psycopg2.connect(**destination.postgres_connect_kwargs())

    try:
        with connection_pg.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    c.table_name,
                    c.column_name,
                    c.data_type,
                    c.is_nullable,
                    c.ordinal_position,
                    pg_catalog.obj_description(
                        (quote_ident(c.table_schema) || '.' || quote_ident(c.table_name))::regclass,
                        'pg_class'
                    ) AS table_description,
                    pg_catalog.col_description(
                        (quote_ident(c.table_schema) || '.' || quote_ident(c.table_name))::regclass,
                        c.ordinal_position
                    ) AS column_description
                FROM information_schema.columns c
                WHERE c.table_schema = %s
                  AND c.table_name NOT LIKE '\\_dlt\\_%%' ESCAPE '\\'
                  AND c.column_name NOT LIKE '\\_dlt\\_%%' ESCAPE '\\'
                ORDER BY c.table_name, c.ordinal_position
                """,
                (destination.schema,),
            )
            rows = cursor.fetchall()
    finally:
        connection_pg.close()

    source_by_table = {item["table_name"]: item for item in loaded_tables}
    tables: dict[str, dict[str, Any]] = {}

    for row in rows:
        table_name = row[0]
        table_source = source_by_table.get(table_name, {})
        table = tables.setdefault(
            table_name,
            {
                "name": table_name,
                "source_name": table_source.get("source_name", table_name),
                "description": row[5] or "",
                "row_count": table_source.get("row_count", 0),
                "columns": [],
                **(
                    {"row_key": table_source["row_key"]}
                    if table_source.get("row_key")
                    else {}
                ),
            },
        )
        source_columns = {
            str(column.get("name")): str(column.get("source_name"))
            for column in table_source.get("columns") or []
            if isinstance(column, dict)
            and column.get("name")
            and column.get("source_name")
        }
        table["columns"].append(
            {
                "name": row[1],
                "type": row[2],
                "nullable": row[3] == "YES",
                "description": row[6] or "",
                **(
                    {"source_name": source_columns[row[1]]}
                    if row[1] in source_columns
                    else {}
                ),
            }
        )

    return {
        "version": 1,
        "connection_id": connection["id"],
        "connection_name": connection["name"],
        "slug": connection["slug"],
        "storage_key": connection["storage_key"],
        "organization_id": connection["organization_id"],
        "plugin": GOOGLE_DRIVE_KEY,
        "source": source_info,
        "destination": {
            "id": destination.id,
            "name": destination.name,
            "slug": destination.slug,
            "type": destination.type,
            "schema": destination.schema,
        },
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "load_ids": load_ids,
        "tables": list(tables.values()),
    }


def _manifest_path(slug: str):
    return config_path(slug).with_suffix(".manifest.yaml")


def _write_manifest(slug: str, manifest: dict[str, Any]) -> None:
    path = _manifest_path(slug)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_suffix(".yaml.tmp")
    temp_path.write_text(
        yaml.safe_dump(manifest, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    temp_path.chmod(0o600)
    temp_path.replace(path)


def _previous_manifest_table_names(slug: str) -> set[str]:
    path = _manifest_path(slug)

    if not path.is_file():
        return set()

    try:
        manifest = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError:
        return set()

    tables = manifest.get("tables") if isinstance(manifest, dict) else []
    return {
        str(table.get("name"))
        for table in tables or []
        if isinstance(table, dict) and table.get("name")
    }


async def _connection(connection_id: int, organization_id: int) -> dict[str, Any]:
    async with db_connection() as db:
        row = await db.fetchrow(
            """
            SELECT c.id, c.name, c.slug, c.storage_key, c.organization_id,
                   c.plugin, c.status, c.created_at,
                   c.last_synced_at, c.last_sync_error,
                   c.destination_id, c.destination_schema,
                   d.name AS destination_name,
                   d.slug AS destination_slug,
                   d.type AS destination_type,
                   d.configuration AS destination_configuration,
                   d.is_builtin AS destination_is_builtin,
                   d.is_default AS destination_is_default
            FROM connections c
            JOIN destinations d ON d.id = c.destination_id
            WHERE c.id = $1 AND c.plugin = $2 AND c.organization_id = $3
            """,
            connection_id,
            GOOGLE_DRIVE_KEY,
            organization_id,
        )

    if not row:
        raise HTTPException(404, "Connection not found")

    return dict(row)


async def _start_run(connection_id: int, trigger: str) -> int:
    async with db_connection() as db, db.transaction():
        run_id = await db.fetchval(
            """
            INSERT INTO sync_runs (connection_id, trigger, status)
            VALUES ($1, $2, 'running')
            RETURNING id
            """,
            connection_id,
            trigger,
        )
        await db.execute(
            """
            UPDATE connections
            SET status = 'syncing', last_sync_started_at = now(),
                last_sync_error = NULL
            WHERE id = $1
            """,
            connection_id,
        )
        return int(run_id)


async def _finish_run(
    run_id: int,
    connection_id: int,
    *,
    result: dict[str, Any] | None = None,
    error: str | None = None,
) -> None:
    result = result or {}
    async with db_connection() as db, db.transaction():
        await db.execute(
            """
            UPDATE sync_runs
            SET status = $1, finished_at = now(), table_count = $2,
                row_count = $3, load_ids = $4, error = $5
            WHERE id = $6
            """,
            "failed" if error else "success",
            result.get("table_count"),
            result.get("row_count"),
            ",".join(result.get("load_ids") or []),
            error,
            run_id,
        )
        await db.execute(
            """
            UPDATE connections
            SET status = $1,
                last_synced_at = CASE WHEN $2::text IS NULL THEN now() ELSE last_synced_at END,
                last_sync_error = $3
            WHERE id = $4
            """,
            "failed" if error else "active",
            error,
            error,
            connection_id,
        )


async def _refresh_models_and_metadata(connection: dict[str, Any]) -> None:
    from app.agent.metadata import get_schema_with_descriptions
    from app.cube.client import CubeAPIError, load_cube_meta
    from app.cube.model import sync_connection_models
    from app.routers.connection_metadata import write_connection_metadata_cache

    destination = runtime_from_connection(connection)
    storage_key = connection.get("storage_key") or connection["slug"]
    schema = await get_schema_with_descriptions(
        destination.schema,
        use_cache=False,
        destination=destination,
        cache_key=storage_key,
    )
    await write_connection_metadata_cache(
        connection_id=connection["id"],
        slug=storage_key,
        display_slug=connection["slug"],
        plugin=connection["plugin"],
        destination_schema=(connection.get("destination_schema") or storage_key),
        live_schema=schema,
    )
    await sync_connection_models()

    try:
        await load_cube_meta()
    except CubeAPIError as exc:
        raise HTTPException(
            502,
            "PostgreSQL sync completed, but Cube rejected the generated "
            f"semantic model: {exc.message}",
        ) from exc


def _google_client_id() -> str:
    value = os.getenv("GOOGLE_OAUTH_CLIENT_ID", "").strip()

    if not value:
        raise ValueError("GOOGLE_OAUTH_CLIENT_ID is not configured")

    return value


def _google_client_secret() -> str:
    value = os.getenv("GOOGLE_OAUTH_CLIENT_SECRET", "").strip()

    if not value:
        raise ValueError("GOOGLE_OAUTH_CLIENT_SECRET is not configured")

    return value
